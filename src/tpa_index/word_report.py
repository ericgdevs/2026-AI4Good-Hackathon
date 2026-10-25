"""006: a plain-English Word report of the newest 005 insights result, for funders and TPA staff (about 3 pages).
Reads result files through demo.InsightsController (no parsing of its own). Aggregates only: no video names, no
clip-level rows, no transcript text, no individual scores; expressions are always unvalidated. Every number comes from
the result files."""
import re
from datetime import datetime, timezone
from pathlib import Path

from tpa_index.config import ConfigError, load_theme
from tpa_index.demo import PHASE_NAMES, InsightsController, indicator_name

FONT = "Calibri"
TITLE = "Program Impact Report"
ORG, PROGRAM = "The Performance Academy", "This Is ME"
CREDIT = "Built for The Performance Academy"
# plain subject for key findings, unit words after a number, one-sentence explanation
PLAIN = {
    "performance_share": ("Time spent performing", "of youth time",
                          "The share of youth time where young people were performing, for example singing, dancing, "
                          "acting or spoken word."),
    "solo_moments_minutes_per_hour": ("Solo performance time", "minutes per hour of youth footage",
                                      "Minutes per hour of youth footage where one young person was performing alone "
                                      "on stage."),
    "youth_voice_minutes_per_hour": ("Youth voice time", "minutes per hour of youth footage",
                                     "Minutes per hour of youth footage where young people performed with their "
                                     "voice: singing, speaking or spoken word."),
    "positive_expression_rate": ("Positive expressions (unvalidated)", "of clips with visible faces",
                                 "Among short clips where faces were visible, the share where the AI suggested a "
                                 "positive expression such as a smile. This is an unchecked suggestion about what the "
                                 "camera shows, never a measure of feelings."),
}
TAG_MEANINGS = {
    "Staff-verified": "Every clip behind this number was checked and confirmed by TPA staff.",
    "Partly verified": "Part of this number was confirmed by staff; the exact tag gives the percentage.",
    "AI estimate (trained)": "No clip behind this number was checked by staff; it comes from an AI model trained on "
                             "staff labels.",
    "AI estimate (unchecked)": "No clip behind this number was checked by staff; it is an AI suggestion only.",
}
EXACT = {"Staff-verified": "staff-verified", "Partly verified": "mixed (NN% staff-verified)",
         "AI estimate (trained)": "model estimate (learned)", "AI estimate (unchecked)": "model estimate (unvalidated)"}
PROTECTION = [
    "Everything runs on one private computer. Nothing is sent to cloud services.",
    "No faces are stored, and no names are recorded.",
    "There are no individual scores. Results describe the footage, never a young person.",
]
GOOD_TO_KNOW = [
    "Numbers marked as AI estimates have not been checked by staff yet. Staff-verified numbers are the most reliable.",
    "Every number describes the footage as a whole, not any individual young person.",
    "The showcase is a performance, so some increase in performing time is expected by design.",
]
SAMPLE_BANNER = "SAMPLE — generated from mock test data. These numbers are not real program results."


def _num(v):
    return float(v) if v not in (None, "") else None


def _ok(r) -> bool:
    return r.get("status", "ok") in ("", "ok") and _num(r["value"]) is not None


def _plain_value(v: float, unit: str) -> str:
    return f"{100 * v:.0f}%" if unit == "share" else f"{v:.1f}"


def key_findings(headline: list[dict]) -> list[str]:
    """One sentence per headline indicator with both an early and a showcase value (at most four)."""
    from tpa_index.ui_html import tag_kind

    out = []
    for ind in dict.fromkeys(r["indicator"] for r in headline):
        by = {r["phase"]: r for r in headline if r["indicator"] == ind}
        early, show = by.get("early"), by.get("showcase")
        if not (early and show and _ok(early) and _ok(show)):
            continue
        a, b = _num(early["value"]), _num(show["value"])
        subject, unit_words, _ = PLAIN.get(ind, (indicator_name(ind), "", ""))
        verb = "rose" if b > a else "fell" if b < a else "stayed the same"
        pa, pb = _plain_value(a, early["unit"]), _plain_value(b, show["unit"])
        la, lb = tag_kind(early["tag"])[0], tag_kind(show["tag"])[0]
        out.append(f"{subject} {verb} from {pa} to {pb} {unit_words}, from early sessions to the showcase "
                   f"({la if la == lb else f'early: {la}; showcase: {lb}'}).".replace("  ", " "))
    return out[:4]


def program_line(ins: InsightsController) -> str:
    """`The Performance Academy · This Is ME 2025`: the year only when the videos share one cohort name with a year."""
    from tpa_index.schema import read_official

    vids = set(ins.state.get("videos", []))
    cohorts = {v["cohort"] for v in read_official("videos", root=ins.root).to_pylist() if v["video_id"] in vids}
    year = re.search(r"\b(19|20)\d{2}\b", cohorts.pop()) if len(cohorts) == 1 and None not in cohorts else None
    return f"{ORG} · {PROGRAM}" + (f" {year.group(0)}" if year else "")


# ---- formatting helpers (python-docx has no boxes, shading or page fields of its own) ----
def _rgb(hex_colour: str):
    from docx.shared import RGBColor

    return RGBColor.from_string(hex_colour.lstrip("#").upper())


def _shade(cell, fill: str):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill.lstrip("#"))
    cell._tc.get_or_add_tcPr().append(shd)


def _borders(table, colour: str):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), colour.lstrip("#"))
        borders.append(el)
    table._tbl.tblPr.append(borders)


def _field(paragraph, code: str):
    """A Word field (PAGE, NUMPAGES) that Word fills in when the document opens or prints."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, code), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {text} "
        run._r.append(el)
    return run


def _space(paragraph, before=0, after=6):
    from docx.shared import Pt

    paragraph.paragraph_format.space_before, paragraph.paragraph_format.space_after = Pt(before), Pt(after)
    return paragraph


def _box(doc, title: str, lines: list[str], theme: dict):
    """A highlighted box: one shaded cell with a bold title and bullets."""
    table = doc.add_table(rows=1, cols=1)
    _borders(table, theme["border"])
    cell = table.cell(0, 0)
    _shade(cell, theme["accent_soft"])
    run = cell.paragraphs[0].add_run(title)
    run.bold, run.font.color.rgb = True, _rgb(theme["accent_deep"])
    _space(cell.paragraphs[0], 4, 4)
    for line in lines:
        _space(cell.add_paragraph(line, style="List Bullet"), 0, 3)
    _space(doc.add_paragraph(), 0, 4)


def _table(doc, header: list[str], rows: list[list[str]], theme: dict):
    """Theme-coloured header row, light shading on every other row, thin borders in the theme border colour.
    A cell's second line (after a newline) is the source label, in smaller secondary text."""
    from docx.shared import Pt

    table = doc.add_table(rows=1, cols=len(header))
    _borders(table, theme["border"])
    for cell, text in zip(table.rows[0].cells, header):
        _shade(cell, theme["accent_deep"])
        run = cell.paragraphs[0].add_run(text)
        run.bold, run.font.color.rgb = True, _rgb(theme["surface"])
    for i, row in enumerate(rows):
        for cell, text in zip(table.add_row().cells, row):
            first, *rest = str(text).split("\n")
            cell.paragraphs[0].add_run(first)
            for extra in rest:
                run = cell.add_paragraph().add_run(extra)
                run.font.size, run.font.color.rgb = Pt(8.5), _rgb(theme["text_secondary"])
            if i % 2:
                _shade(cell, theme["background"])
    _space(doc.add_paragraph(), 0, 2)
    return table


def _heading(doc, text: str):
    return _space(doc.add_heading(text, level=1), 14, 6)


def write_word_report(dest, root=None, sample: bool = False) -> Path:
    """Write the report for the newest complete insights folder (and evaluate folder, if any) to dest.
    sample=True marks a report built from mock data: a banner under the title and in the footer."""
    import docx
    from docx.oxml.ns import qn
    from docx.shared import Pt

    from tpa_index.report import show
    from tpa_index.ui_html import tag_kind

    ins = InsightsController(root)
    if not ins.available:
        raise ConfigError(ins.message())
    st, theme = ins.state, load_theme()
    doc = docx.Document()
    props = doc.core_properties  # the template's defaults name the library; replace them
    props.author, props.last_modified_by, props.comments = "TPA video index", "TPA video index", ""
    props.title = TITLE
    for name in ("Normal", "Title", "Heading 1", "List Bullet"):
        style = doc.styles[name]
        style.font.name = FONT
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    doc.styles["Normal"].font.size, doc.styles["Normal"].font.color.rgb = Pt(10.5), _rgb(theme["text"])
    doc.styles["Title"].font.size, doc.styles["Heading 1"].font.size = Pt(26), Pt(14)
    for name in ("Title", "Heading 1"):
        doc.styles[name].font.color.rgb = _rgb(theme["accent_deep"])

    # (1) cover block
    _space(doc.add_heading(TITLE, level=0), 0, 2)
    sub = doc.add_paragraph().add_run(program_line(ins))
    sub.bold, sub.font.size = True, Pt(13)
    doc.add_paragraph().add_run(f"Generated {datetime.now(timezone.utc):%d %B %Y}").font.color.rgb = \
        _rgb(theme["text_secondary"])
    if sample:
        banner = doc.add_paragraph().add_run(SAMPLE_BANNER)
        banner.bold, banner.font.size, banner.font.color.rgb = True, Pt(12), _rgb(theme["accent_deep"])

    # (2) at a glance
    head = st["headline"]
    _box(doc, "At a glance", key_findings(head) or ["No early-to-showcase comparison is available in these results."],
         theme)

    # (3) youth indicators by phase: each value with its short source label; the legend gives the exact tags
    _heading(doc, "Youth indicators by phase")
    phases = [p for p in ("early", "mid", "showcase") if any(r["phase"] == p for r in head)]
    rows, tags = [], []
    for ind in dict.fromkeys(r["indicator"] for r in head):
        by = {r["phase"]: r for r in head if r["indicator"] == ind}
        cells = []
        for p in phases:
            r = by.get(p)
            if r and _ok(r):
                cells.append(f"{show(_num(r['value']), r['unit'])}\n{tag_kind(r['tag'])[0]}")
                tags.append(r["tag"])
            else:
                cells.append(r["status"] if r else "")
        rows.append([PLAIN.get(ind, (indicator_name(ind),))[0], *cells])
    _table(doc, ["Measure", *[PHASE_NAMES[p] for p in phases]], rows, theme)
    legend = doc.add_paragraph()
    legend.add_run("Source labels: ").bold = True
    grouped: dict[str, list[str]] = {}
    for t in dict.fromkeys(tags):  # each short label once, with every exact tag it stands for here
        grouped.setdefault(tag_kind(t)[0], []).append(t)
    legend.add_run("; ".join(f"{label} = {', '.join(ts)}" for label, ts in grouped.items()) + ".")
    for run in legend.runs:
        run.font.size, run.font.color.rgb = Pt(8.5), _rgb(theme["text_secondary"])
    for ind in dict.fromkeys(r["indicator"] for r in head):
        if ind in PLAIN:
            p = _space(doc.add_paragraph(style="List Bullet"), 0, 2)
            p.add_run(PLAIN[ind][0] + ": ").bold = True
            p.add_run(PLAIN[ind][2])

    # (4) who was leading
    _heading(doc, "Who was leading?")
    t = st["time"]
    if t:
        parts = [("Young people leading", "youth_s"), ("Staff leading, confirmed by staff", "staff_led_verified_s"),
                 ("Staff leading, AI estimate", "staff_led_model_s"), ("Idle or setup", "idle_s")]
        total = _num(t["total_s"]) or 0.0
        _table(doc, ["Time", "Minutes", "Share of all footage"],
               [[name, f"{_num(t[k]) / 60:.1f}", f"{100 * _num(t[k]) / total:.0f}%" if total else "n/a"]
                for name, k in parts], theme)
        doc.add_paragraph("Staff-led and idle time are left out of every youth number in this report.")
    else:
        doc.add_paragraph("Not available in these results.")

    # (5) how we measured it
    _heading(doc, "How we measured it")
    n_videos, n_clips = len(st.get("videos", [])), st.get("clips")
    clips = f" and {n_clips} short clips" if n_clips is not None else ""
    doc.add_paragraph(f"We looked at {n_videos} video{'s' if n_videos != 1 else ''}{clips}. The AI finds moments in "
                      "the footage and suggests what is happening; TPA staff then check and confirm those suggestions.")
    r = ins.agreement_row()
    v = _num(r["primary_accuracy"]) if r else None
    if v is None:
        doc.add_paragraph("Accuracy check: not available yet. No comparison with staff labels has been run.")
    else:
        lo, hi = _num(r["primary_accuracy_ci_low"]), _num(r["primary_accuracy_ci_high"])
        m = round(_num(r["n_segments"]) or 0)
        rng = f" (likely range {100 * lo:.0f}–{100 * hi:.0f}%)" if lo is not None and hi is not None else ""
        doc.add_paragraph(f"Accuracy check: a staff member labelled {m} clips without seeing the AI's suggestions. "
                          f"The AI's main activity matched the staff label in {100 * v:.0f}% of them{rng}. AI "
                          "suggestions are a starting point; staff review confirms them.")

    # (6) how the data is protected
    _heading(doc, "How the data is protected")
    for line in PROTECTION:
        _space(doc.add_paragraph(line, style="List Bullet"), 0, 2)

    # (7) what the labels mean
    _box(doc, "What the labels mean",
         [f"{label}: {meaning} (tag: {EXACT[label]})" for label, meaning in TAG_MEANINGS.items()], theme)

    # (8) good to know
    _heading(doc, "Good to know")
    for line in GOOD_TO_KNOW:
        _space(doc.add_paragraph(line, style="List Bullet"), 0, 2)

    # (9) footer: credit and page numbers (plus the sample banner when built from mock data)
    footer = doc.sections[0].footer.paragraphs[0]
    footer.add_run(f"{CREDIT} · Page ")
    _field(footer, "PAGE")
    footer.add_run(" of ")
    _field(footer, "NUMPAGES")
    for run in footer.runs:
        run.font.size, run.font.color.rgb = Pt(8), _rgb(theme["text_secondary"])
    if sample:
        footer.insert_paragraph_before().add_run(SAMPLE_BANNER).bold = True

    dest = Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(dest)
    return dest
