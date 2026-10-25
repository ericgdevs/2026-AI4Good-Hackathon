"""006: plain-English Word report of the insights results, on 005's mock world plus a real 005 insights run."""
import re

import docx
import pytest

from tests.insights_fixtures import SECRET, SENTINEL, build_world
from tests.test_insights import _evaluate_folder, _rows
from tpa_index import cli, report, word_report
from tpa_index.demo import InsightsController
from tpa_index.schema import read_official
from tpa_index.config import ConfigError
from tpa_index.indicators import LIMITATIONS
from tpa_index.ui_html import tag_kind

SECTIONS = ["Program Impact Report", "At a glance", "Youth indicators by phase", "Who was leading?",
            "How we measured it", "How the data is protected", "What the labels mean", "Good to know"]
JARGON = ["F1", "kappa", "macro", "zero-shot", "ablation", "threshold", "SigLIP", "Whisper", "YOLO", "segment",
          "transcript", "AST", "EmotiEff", "logistic"]


def document_text(path) -> str:
    """Body paragraphs and table cells in document order, then the footer and file properties."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(path)
    parts = []
    for block in d.iter_inner_content():
        if isinstance(block, Paragraph):
            parts.append(block.text)
        elif isinstance(block, Table):
            parts += [p.text for row in block.rows for c in row.cells for p in c.paragraphs]
    parts += [p.text for s in d.sections for p in s.footer.paragraphs]
    props = d.core_properties
    parts += [props.author, props.last_modified_by, props.comments, props.title]
    return "\n".join(parts)


@pytest.fixture
def world(tmp_path, monkeypatch):
    w = build_world(tmp_path, monkeypatch)
    assert cli.main(["insights"]) == 0
    return {**w, "folder": report.latest_result("insights")}


def test_report_sections_tags_and_findings(world, tmp_path):
    path = word_report.write_word_report(tmp_path / "r.docx")
    text = document_text(path)
    positions = [text.index(s) for s in SECTIONS]
    assert positions == sorted(positions)
    assert word_report.program_line(InsightsController()) in text
    head = _rows(world["folder"] / "headline.csv")
    for r in head:  # value, then the Scorecard's short source label; the legend gives every exact tag
        assert f"{report.show(float(r['value']), r['unit'])}\n{tag_kind(r['tag'])[0]}" in text
        legend = next(line for line in text.splitlines() if line.startswith("Source labels:"))
        assert r["tag"] in legend and legend.count(f"{tag_kind(r['tag'])[0]} = ") == 1
    assert "Expression Index" not in text and report.INDEX_CAVEAT not in text
    assert not any(line in text for line in LIMITATIONS)  # replaced by the three plain "Good to know" bullets
    assert all(line in text for line in word_report.GOOD_TO_KNOW) and len(word_report.GOOD_TO_KNOW) == 3
    assert all(line in text for line in word_report.PROTECTION)
    assert "python-docx" not in text and "Accuracy check: not available yet" in text  # no evaluate folder yet
    clips = read_official("segments").num_rows
    assert f"We looked at 4 videos and {clips} short clips." in text

    findings = word_report.key_findings(head)
    assert 3 <= len(findings) <= 4 and all(f in text for f in findings)
    by = {(r["indicator"], r["phase"]): r for r in head}
    voice = (by[("youth_voice_minutes_per_hour", "early")], by[("youth_voice_minutes_per_hour", "showcase")])
    a, b = (float(r["value"]) for r in voice)
    expected = (f"Youth voice time {'rose' if b > a else 'fell' if b < a else 'stayed the same'} from {a:.1f} to "
                f"{b:.1f} minutes per hour of youth footage")
    assert any(f.startswith(expected) for f in findings)


def test_report_formatting(world, tmp_path):
    """Theme colours on headings and the table header, a standard font, page numbers and the credit in the footer."""
    from tpa_index.config import load_theme

    d = docx.Document(word_report.write_word_report(tmp_path / "r.docx"))
    theme = load_theme()
    assert d.styles["Normal"].font.name == word_report.FONT == "Calibri"
    assert str(d.styles["Heading 1"].font.color.rgb) == theme["accent_deep"].lstrip("#").upper()
    table = d.tables[1]  # indicators (the first table is the At a glance box)
    header_xml = table.rows[0].cells[0]._tc.xml
    assert theme["accent_deep"].lstrip("#").upper() in header_xml.upper()
    assert theme["background"].lstrip("#").upper() in table.rows[2].cells[0]._tc.xml.upper()  # light row shading
    footer = d.sections[0].footer.paragraphs[-1]
    assert footer.text.startswith(word_report.CREDIT) and " PAGE " in footer._p.xml and " NUMPAGES " in footer._p.xml


def test_report_with_evaluate_and_plain_words(world, tmp_path):
    _evaluate_folder(world["folder"].parent)
    text = document_text(word_report.write_word_report(tmp_path / "r.docx"))
    assert ("a reviewer labelled 30 clips without seeing the AI's suggestions. The AI's main activity matched the "
            "blind human label in 70% of them (likely range 50–85%).") in text
    for word in JARGON:
        assert not re.search(rf"\b{re.escape(word)}", text, re.IGNORECASE), word
    assert SENTINEL not in text and SECRET not in text


def test_report_needs_insights(tmp_path, monkeypatch):
    build_world(tmp_path, monkeypatch)
    with pytest.raises(ConfigError, match="not available"):
        word_report.write_word_report(tmp_path / "r.docx")


def test_key_findings_handle_missing_phases():
    rows = [{"indicator": "performance_share", "phase": "early", "value": "0.4", "unit": "share", "tag": "staff-verified",
             "status": "ok"},
            {"indicator": "performance_share", "phase": "showcase", "value": "0.65", "unit": "share",
             "tag": "mixed (40% staff-verified)", "status": "ok"},
            {"indicator": "solo_moments_minutes_per_hour", "phase": "showcase", "value": "", "unit": "",
             "tag": "", "status": "no videos in this phase"}]
    assert word_report.key_findings(rows) == [
        "Time spent performing rose from 40% to 65% of youth time, from early sessions to the showcase "
        "(early: Staff-verified; showcase: Partly verified)."]


def test_sample_flag_adds_banner(world, tmp_path):
    from docx import Document

    normal = document_text(word_report.write_word_report(tmp_path / "r.docx"))
    assert word_report.SAMPLE_BANNER not in normal and "SAMPLE" not in normal
    path = word_report.write_word_report(tmp_path / "s.docx", sample=True)
    d = Document(path)
    assert d.paragraphs[0].text == word_report.TITLE
    assert word_report.SAMPLE_BANNER in [p.text for p in d.paragraphs[:4]]  # in the cover block, under the title
    assert word_report.SAMPLE_BANNER in [p.text for p in d.sections[0].footer.paragraphs]
